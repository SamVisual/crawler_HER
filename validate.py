import argparse
import logging
import re
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from random import uniform
from typing import List, Tuple, Optional

import requests
from bs4 import BeautifulSoup
from fake_useragent import FakeUserAgent
from Levenshtein import ratio
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from tqdm import tqdm
from urllib.parse import quote_plus
from urllib import robotparser


logging.basicConfig(
    filename="validate.log",
    level=logging.INFO,
    format="%(asctime)s %(levelname)s: %(message)s",
)


HEADERS_LIST = [
    {
        "User-Agent": FakeUserAgent().random,
        "Accept-Language": "de-DE,de;q=0.9",
    }
    for _ in range(10)
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Validate Excel records via web search")
    parser.add_argument("--file", required=True, help="Excel file to process")
    parser.add_argument(
        "--delay",
        default="2-5",
        help="Random delay range in seconds, e.g. 2-5",
    )
    parser.add_argument(
        "--threads",
        type=int,
        default=1,
        help="Number of worker threads (default 1)",
    )
    return parser.parse_args()


def build_query(row_values: Tuple[str, str, Optional[str]]) -> str:
    """Build search query from (name, address, phone)."""
    name, address, phone = row_values
    parts = [name, address]
    if phone:
        parts.append(str(phone))
    return " ".join(str(p) for p in parts if p)


def allowed_by_robots(url: str, user_agent: str) -> bool:
    rp = robotparser.RobotFileParser()
    rp.set_url("https://duckduckgo.com/robots.txt")
    try:
        rp.read()
    except Exception as exc:
        logging.warning("Failed to read robots.txt: %s", exc)
        return False
    return rp.can_fetch(user_agent, url)


def search_web(query: str, session: requests.Session) -> List[Tuple[str, str, str]]:
    """Search DuckDuckGo and return list of (title, url, snippet)."""
    headers = session.headers
    search_url = f"https://html.duckduckgo.com/html/?q={quote_plus(query)}"
    if not allowed_by_robots("/html/", headers.get("User-Agent", "")):
        logging.warning("Search disallowed by robots.txt")
        return []
    resp = session.get(search_url, timeout=10)
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")
    results = []
    for item in soup.select("div.result")[:5]:
        a = item.select_one("a.result__a[href]")
        if not a:
            continue
        url = a["href"]
        title = a.get_text(strip=True)
        snippet_tag = item.select_one("span.result__snippet")
        snippet = snippet_tag.get_text(strip=True) if snippet_tag else ""
        results.append((title, url, snippet))
    return results


def evaluate_results(
    results: List[Tuple[str, str, str]],
    reference: Tuple[str, str, Optional[str]],
) -> Tuple[bool, List[str], Optional[str]]:
    name, address, phone = reference
    postal_ref = None
    postal_match = re.search(r"\b(\d{5})\b", address or "")
    if postal_match:
        postal_ref = postal_match.group(1)

    def matches(res: Tuple[str, str, str]) -> bool:
        title, url, snippet = res
        if postal_ref and postal_ref not in snippet:
            return False
        if phone:
            digits = re.sub(r"\D", "", str(phone))
            if digits and digits not in re.sub(r"\D", "", snippet):
                return False
        similarity = ratio(name.lower(), title.lower())
        return similarity >= 0.8

    urls = [u for _, u, _ in results]
    for res in results:
        if matches(res):
            return True, urls, res[1]
    return False, urls, urls[0] if urls else None


def auto_fit_columns(sheet):
    for column_cells in sheet.columns:
        length = max(len(str(cell.value)) if cell.value else 0 for cell in column_cells)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = length + 2


def process_row(idx: int, row, headers: List[str], delay_range: Tuple[float, float]) -> Tuple[int, bool, List[str], Optional[str]]:
    name = row[headers.index("Name")].value
    address = row[headers.index("Adresse")].value
    phone = row[headers.index("Telefon")].value if "Telefon" in headers else None
    if not name and not address:
        return idx, False, [], None
    query = build_query((name, address, phone))
    session = requests.Session()
    session.headers.update(HEADERS_LIST[idx % len(HEADERS_LIST)])
    results = []
    try:
        results = search_web(query, session)
    except Exception as exc:
        logging.error("Search error for row %s: %s", idx, exc)
    found, urls, top = evaluate_results(results, (name, address, phone))
    delay = uniform(*delay_range)
    time.sleep(delay)
    return idx, found, urls, top


def main():
    args = parse_args()
    delay_bounds = tuple(float(x) for x in args.delay.split("-"))
    wb: Workbook = load_workbook(args.file)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]

    # Append new headers if not present
    new_headers = ["Gefunden?", "Kandidaten-URLs", "Erster_Treffer", "Timestamp"]
    for h in new_headers:
        if h not in headers:
            headers.append(h)
            sheet.cell(row=1, column=len(headers)).value = h

    rows = list(sheet.iter_rows(min_row=2, max_col=len(headers), values_only=False))

    results = [None] * len(rows)
    with ThreadPoolExecutor(max_workers=args.threads) as executor:
        futures = []
        for idx, row in enumerate(rows, start=2):
            futures.append(
                executor.submit(process_row, idx, row, headers, delay_bounds)
            )
        for fut in tqdm(futures, desc="Validating", unit="row"):
            idx, found, urls, top = fut.result()
            row = sheet[idx]
            base = len(headers) - len(new_headers)
            row[base].value = "✔" if found else "✘"
            row[base + 1].value = ",".join(urls)
            row[base + 2].value = top or ""
            row[base + 3].value = datetime.utcnow().isoformat()

    sheet.freeze_panes = "A2"
    auto_fit_columns(sheet)
    out_file = args.file.replace(".xlsx", "_validated.xlsx")
    wb.save(out_file)
    logging.info("Saved validated workbook to %s", out_file)


if __name__ == "__main__":
    main()
